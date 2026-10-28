from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

PRODUCTS = Path("data/seed/products.csv")
OUTPUT = Path("templates/offline_sales_template.xlsx")
LAST_ROW = 2001

HEADERS = [
    "receipt_number",
    "sale_date",
    "sku",
    "product_name_auto",
    "quantity",
    "unit_price_jpy",
    "customer_email",
]

INSTRUCTIONS = [
    "Offline sales template",
    "",
    "One row per product sold. A receipt with several products uses several rows with the same receipt_number.",
    "receipt_number: required, must be unique for each sale (for example R-2026-0001).",
    "sale_date: required, the day of the sale (yyyy-mm-dd).",
    "sku: required, choose from the dropdown. The product name appears next to it so you can check it.",
    "quantity: required, a whole number of 1 or more.",
    "unit_price_jpy: optional. Leave empty to use the catalog price. Fill it in when you gave a discount.",
    "customer_email: optional. Leave empty for walk-in customers.",
    "",
    "Do not rename the columns or the sheets. Do not change column D (product_name_auto).",
    "Rows that fail the checks are not loaded. They are listed with a reason in the load report.",
]


def add_validation(sheet, validation, cell_range):
    sheet.add_data_validation(validation)
    validation.add(cell_range)


def main():
    products = pd.read_csv(PRODUCTS)
    workbook = Workbook()

    sales = workbook.active
    sales.title = "Sales"
    sales.append(HEADERS)
    for cell in sales[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    sales.freeze_panes = "A2"
    for column, width in zip("ABCDEFG", (18, 14, 14, 44, 10, 16, 30)):
        sales.column_dimensions[column].width = width

    for row in range(2, LAST_ROW + 1):
        sales.cell(row=row, column=2).number_format = "yyyy-mm-dd"
        sales.cell(
            row=row,
            column=4,
            value=(
                f'=IF(C{row}="","",IFERROR(INDEX(Products!$B:$B,'
                f'MATCH(C{row},Products!$A:$A,0)),"unknown SKU"))'
            ),
        )

    last_product_row = len(products) + 1
    add_validation(
        sales,
        DataValidation(
            type="list",
            formula1=f"=Products!$A$2:$A${last_product_row}",
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="Unknown SKU",
            error="Choose a SKU from the dropdown list (see the Products sheet).",
        ),
        f"C2:C{LAST_ROW}",
    )
    add_validation(
        sales,
        DataValidation(
            type="whole",
            operator="greaterThanOrEqual",
            formula1="1",
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="Quantity",
            error="Quantity must be a whole number of 1 or more.",
        ),
        f"E2:E{LAST_ROW}",
    )
    add_validation(
        sales,
        DataValidation(
            type="date",
            operator="greaterThanOrEqual",
            formula1="42005",
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="Date",
            error="Enter a valid date, for example 2026-10-25.",
        ),
        f"B2:B{LAST_ROW}",
    )
    add_validation(
        sales,
        DataValidation(
            type="decimal",
            operator="greaterThanOrEqual",
            formula1="0",
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="Price",
            error="Price must be a number of 0 or more.",
        ),
        f"F2:F{LAST_ROW}",
    )

    sheet = workbook.create_sheet("Products")
    sheet.append(["sku", "name", "maker", "category", "price_jpy"])
    for row in products[["sku", "name", "maker", "category", "price_jpy"]].itertuples(index=False):
        sheet.append(list(row))
    sheet.column_dimensions["A"].width = 12
    sheet.column_dimensions["B"].width = 50

    guide = workbook.create_sheet("Instructions")
    for line in INSTRUCTIONS:
        guide.append([line])
    guide.column_dimensions["A"].width = 110
    guide["A1"].font = Font(bold=True, size=14)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)
    print("saved", OUTPUT)


if __name__ == "__main__":
    main()